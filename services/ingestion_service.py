"""
Ingestion Service

Pipeline for processing uploaded documents:
  1. Download file from Supabase Storage
  2. Extract text (PDF, Excel, DOCX, TXT)
  3. Split into chunks
  4. Embed chunks in batch
  5. Insert into document_chunks (pgvector)
  6. Update document embedding_status

Runs as a FastAPI BackgroundTask — called after the upload endpoint
returns 201. If the process restarts mid-ingestion, the document stays
in 'processing' status. A future admin endpoint can re-trigger ingestion
for stuck documents.
"""

import io
import logging
from typing import Optional

from core.supabase import get_supabase
from services.embedding_service import get_embeddings_batch

logger = logging.getLogger(__name__)

# ── Chunking config ───────────────────────────────────────────────────────────
# 800 tokens ≈ ~600 words. Overlap ensures context isn't lost at chunk boundaries.
CHUNK_SIZE = 800        # tokens (approximate — we split by chars, ~4 chars/token)
CHUNK_OVERLAP = 100     # tokens
CHARS_PER_TOKEN = 4

CHUNK_SIZE_CHARS = CHUNK_SIZE * CHARS_PER_TOKEN        # 3200
CHUNK_OVERLAP_CHARS = CHUNK_OVERLAP * CHARS_PER_TOKEN  # 400


# ── Text extraction ───────────────────────────────────────────────────────────

def _extract_text_from_pdf(data: bytes) -> str:
    """Extract full text from a PDF using PyMuPDF."""
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    pages = []
    for page in doc:
        pages.append(page.get_text())
    doc.close()
    return "\n\n".join(pages)


def _extract_text_from_docx(data: bytes) -> str:
    """Extract text from a .docx file."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_text_from_excel(data: bytes) -> str:
    """
    Extract text from an Excel file.
    Formats each row as 'col1: val1 | col2: val2 | ...' so the
    embedding captures both field names and values semantically.
    """
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows_text = []

    for sheet in wb.worksheets:
        headers = []
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            # Skip fully empty rows
            if all(cell is None for cell in row):
                continue

            if row_idx == 0:
                # Treat first row as headers
                headers = [str(c) if c is not None else f"col{i}" for i, c in enumerate(row)]
            else:
                parts = []
                for header, cell in zip(headers, row):
                    if cell is not None and str(cell).strip():
                        parts.append(f"{header}: {cell}")
                if parts:
                    rows_text.append(" | ".join(parts))

    wb.close()
    return "\n".join(rows_text)


def _extract_text_from_txt(data: bytes) -> str:
    """Decode plain text, trying UTF-8 then latin-1 as fallback."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def extract_text(data: bytes, mime_type: str, filename: str) -> str:
    """
    Route to the correct extractor based on MIME type or filename extension.
    Returns extracted text. Raises ValueError for unsupported formats.
    """
    mime_type = (mime_type or "").lower()
    filename = (filename or "").lower()

    if mime_type == "application/pdf" or filename.endswith(".pdf"):
        return _extract_text_from_pdf(data)

    if mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ) or filename.endswith(".docx") or filename.endswith(".doc"):
        return _extract_text_from_docx(data)

    if mime_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ) or filename.endswith(".xlsx") or filename.endswith(".xls"):
        return _extract_text_from_excel(data)

    if mime_type.startswith("text/") or filename.endswith(".txt") or filename.endswith(".md"):
        return _extract_text_from_txt(data)

    raise ValueError(f"Unsupported file type: mime='{mime_type}' filename='{filename}'")


# ── Chunking ──────────────────────────────────────────────────────────────────

def chunk_text(text: str) -> list[str]:
    """
    Split text into overlapping chunks by character count.
    Prefers splitting on paragraph boundaries (\n\n) when possible.
    Falls back to hard character splits.
    """
    text = text.strip()
    if not text:
        return []

    # Split on double newlines (paragraph boundaries) first
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    chunks: list[str] = []
    current = ""

    for para in paragraphs:
        # If adding this paragraph exceeds chunk size, flush current and start new
        if len(current) + len(para) > CHUNK_SIZE_CHARS and current:
            chunks.append(current.strip())
            # Keep overlap: take the last CHUNK_OVERLAP_CHARS of current chunk
            current = current[-CHUNK_OVERLAP_CHARS:] + "\n\n" + para
        else:
            current = (current + "\n\n" + para) if current else para

        # Hard split if a single paragraph is larger than chunk size
        while len(current) > CHUNK_SIZE_CHARS:
            chunks.append(current[:CHUNK_SIZE_CHARS].strip())
            current = current[CHUNK_SIZE_CHARS - CHUNK_OVERLAP_CHARS:]

    if current.strip():
        chunks.append(current.strip())

    return chunks


# ── Main ingestion pipeline ───────────────────────────────────────────────────

async def ingest_document(document_id: str) -> None:
    """
    Full ingestion pipeline for a single document.
    Called as a FastAPI BackgroundTask after upload.

    Updates embedding_status:
      'processing' → while running
      'done'       → on success
      'error'      → on failure (error detail stored in metadata)
    """
    supabase = get_supabase()

    # Mark as processing
    supabase.table("documents").update(
        {"embedding_status": "processing"}
    ).eq("id", document_id).execute()

    try:
        # 1. Fetch document metadata
        doc_result = (
            supabase.table("documents")
            .select("*")
            .eq("id", document_id)
            .single()
            .execute()
        )
        if not doc_result.data:
            raise ValueError(f"Document {document_id} not found")

        doc = doc_result.data
        storage_path: str = doc["storage_path"]
        mime_type: str = doc.get("mime_type", "")
        filename: str = doc.get("name", "")

        # 2. Download file from Supabase Storage
        file_bytes = supabase.storage.from_("documents").download(storage_path)

        # 3. Extract text
        text = extract_text(file_bytes, mime_type, filename)
        if not text.strip():
            raise ValueError("No text could be extracted from this document")

        # 4. Chunk
        chunks = chunk_text(text)
        if not chunks:
            raise ValueError("Document produced no text chunks after processing")

        logger.info(
            f"Document {document_id} ({filename}): {len(chunks)} chunks from "
            f"{len(text)} chars"
        )

        # 5. Embed all chunks in batch
        embeddings = await get_embeddings_batch(chunks)

        # 6. Delete any existing chunks (re-ingestion case)
        supabase.table("document_chunks").delete().eq(
            "document_id", document_id
        ).execute()

        # 7. Insert chunks + embeddings
        # Supabase Python client handles vector columns as lists of floats.
        rows = [
            {
                "document_id": document_id,
                "chunk_index": idx,
                "content": chunk,
                "embedding": embedding,
                "metadata": {"chunk_index": idx, "total_chunks": len(chunks)},
            }
            for idx, (chunk, embedding) in enumerate(zip(chunks, embeddings))
        ]

        # Insert in batches of 50 to avoid request size limits
        BATCH_SIZE = 50
        for i in range(0, len(rows), BATCH_SIZE):
            supabase.table("document_chunks").insert(rows[i : i + BATCH_SIZE]).execute()

        # 8. Mark as done
        supabase.table("documents").update(
            {"embedding_status": "done"}
        ).eq("id", document_id).execute()

        logger.info(f"Document {document_id} ingested successfully ({len(rows)} chunks)")

    except Exception as e:
        logger.error(f"Ingestion failed for document {document_id}: {e}")
        supabase.table("documents").update({
            "embedding_status": "error",
            "metadata": {"ingestion_error": str(e)},
        }).eq("id", document_id).execute()


async def ingest_phase_output(
    project_id: str,
    phase_id: str,
    run_id: str,
    output_payload: dict,
    created_by: str,
) -> None:
    """
    Auto-ingest the output of a completed phase run as a RAG document.
    Called from routers/webhooks.py after a successful callback.

    The output is stored as a JSON text blob so future phases can
    retrieve relevant context from past executions.
    """
    import json

    supabase = get_supabase()

    # Convert output payload to searchable text
    text = json.dumps(output_payload, ensure_ascii=False, indent=2)
    filename = f"{phase_id}_run_{run_id[:8]}.json"

    # Upload the JSON text to Storage so there's a retrievable artifact
    storage_path = f"phase_outputs/{project_id}/{phase_id}/{run_id}.json"
    supabase.storage.from_("documents").upload(
        storage_path,
        text.encode("utf-8"),
        {"content-type": "application/json"},
    )

    # Create document record
    doc_result = supabase.table("documents").insert({
        "name": filename,
        "type": "project_output",
        "source": f"phase:{phase_id}",
        "project_id": project_id,
        "storage_path": storage_path,
        "mime_type": "application/json",
        "file_size_bytes": len(text.encode("utf-8")),
        "embedding_status": "processing",
        "metadata": {"phase_id": phase_id, "run_id": run_id},
        "created_by": created_by,
    }).execute()

    if not doc_result.data:
        logger.error(f"Failed to create document record for phase output {run_id}")
        return

    document_id = doc_result.data[0]["id"]

    # Chunk and embed directly (no file download needed — we have the text)
    try:
        chunks = chunk_text(text)
        if not chunks:
            return

        embeddings = await get_embeddings_batch(chunks)

        rows = [
            {
                "document_id": document_id,
                "chunk_index": idx,
                "content": chunk,
                "embedding": embedding,
                "metadata": {
                    "phase_id": phase_id,
                    "run_id": run_id,
                    "chunk_index": idx,
                },
            }
            for idx, (chunk, embedding) in enumerate(zip(chunks, embeddings))
        ]

        BATCH_SIZE = 50
        for i in range(0, len(rows), BATCH_SIZE):
            supabase.table("document_chunks").insert(rows[i : i + BATCH_SIZE]).execute()

        supabase.table("documents").update(
            {"embedding_status": "done"}
        ).eq("id", document_id).execute()

        logger.info(
            f"Phase output ingested: phase={phase_id} run={run_id[:8]} "
            f"chunks={len(rows)}"
        )

    except Exception as e:
        logger.error(f"Failed to ingest phase output {run_id}: {e}")
        supabase.table("documents").update(
            {"embedding_status": "error"}
        ).eq("id", document_id).execute()